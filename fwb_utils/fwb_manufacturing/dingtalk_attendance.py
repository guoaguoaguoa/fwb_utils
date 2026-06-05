# Copyright (c) 2026, WenZhou Furui Handicraft Co.,Ltd. and Contributors
# See license.txt
"""钉钉「月度汇总」导入 + 加班自动结算（设计文档 §7.6.2）。

加班规则（用户 2026-06 确认，仅「普工结构」适用）：
- 起算点、起步分钟、阶梯分钟、封顶小时、折天分母和工资日分母均读「薪资考勤参数」。
- 默认值仍保持旧口径：17:00 起算 / 1 小时起步 / 30 分钟阶梯 / 封顶 5 小时 / 7 小时 = 1 天 / 月固定÷30。
- **非普工**：不计加班费（钉钉端按调休处理），overtime_days 一律 0。

纯逻辑（parse_day_cell / overtime_days_for_regular / daily_overtime_days /
parse_monthly_summary）不依赖 DB，可沙箱单测。
"""
from __future__ import annotations

import re

import frappe
from frappe.utils import add_days, cint, date_diff, flt, getdate

DEFAULT_OT_START = "17:00"  # 标准下班，加班起算点
DEFAULT_OT_MIN_MINUTES = 60  # 1 小时起步
DEFAULT_OT_STEP_MIN = 30  # 30 分钟一个阶梯
DEFAULT_OT_CAP_HOURS = 5.0  # 封顶 5 小时
DEFAULT_OT_HOURS_PER_DAY = 7.0  # 7 小时 = 1 天（3.5h=0.5天）
DEFAULT_OVERTIME_DAILY_DIVISOR = 30  # 加班工资 = 折算加班天 × 月固定 / 本数
OVERTIME_COMPONENT = "加班"
MEAL_COMPONENT = "餐补"
LATE_DEDUCTION_COMPONENT = "迟到扣款"
EARLY_DEDUCTION_COMPONENT = "早退扣款"
PERSONAL_SOCIAL_SECURITY_COMPONENT = "社保-个人"
COMPANY_SOCIAL_SECURITY_COMPONENT = "社保-单位"
SOCIAL_SECURITY_EMPLOYMENT_TYPE = "缴纳社保"
DEFAULT_PAID_LEAVE_NAMES = ("年假", "丧假")
REGULAR_WORKER_STRUCTURE = "普工结构"  # 仅此结构计加班费
NON_WORKER_STRUCTURES = ("内贸业务员结构", "外贸业务员结构", "美工结构", "行政结构")
NON_WORKER_MONTHLY_REST_DAYS = 6  # 默认；实际从「薪资考勤参数」读取
NON_WORKER_STANDARD_HOURS_PER_DAY = 7.5  # 默认非普工日工时（08:30-17:30 撇 1.5h 午休）
ATTENDANCE_DEDUCTION_THRESHOLD_MINUTES = 30  # 默认；实际从「薪资考勤参数」读取
DEFAULT_MEAL_UNIT_PRICE = 14.0
DEFAULT_REGULAR_DAILY_DIVISOR = 30  # 普工单日工资 = 月固定 / 本数
FIXED_ATTENDANCE_COMPONENTS = (
	("底薪", "base"),
	("工龄补贴", "custom_seniority_base"),
	("等级工资", "custom_grade_base"),
	("证书津贴", "custom_cert_base"),
)
# 普工月固定 = 底薪+工龄+等级（无证书、无餐补）
REGULAR_FIXED_COMPONENTS = (
	("底薪", "base"),
	("工龄补贴", "custom_seniority_base"),
	("等级工资", "custom_grade_base"),
)
# 排班档案默认值：普工 8.5h(4卡) / 综合岗 7.5h；午休时段不同。实际从「薪资考勤参数·排班档案」读取。
DEFAULT_SCHEDULE_PROFILES = {
	"regular": {"work_start": "07:30", "lunch_start": "11:30", "lunch_end": "12:30", "work_end": "17:00", "daily_hours": 8.5},
	"non_worker": {"work_start": "08:30", "lunch_start": "11:30", "lunch_end": "13:00", "work_end": "17:30", "daily_hours": 7.5},
}


def _is_param_blank(value):
	return value is None or str(value).strip() == ""


def _param_value(doc, fieldname, default, cast=None):
	if not doc:
		return default
	value = doc.get(fieldname)
	if _is_param_blank(value):
		return default
	return cast(value) if cast else value


def _to_minutes(hhmm: str) -> int:
	h, m = hhmm.split(":")
	return int(h) * 60 + int(m)


def _minutes_from_time_value(value, default=None):
	if value in (None, ""):
		value = default
	if value in (None, ""):
		return None
	if hasattr(value, "hour") and hasattr(value, "minute"):
		return cint(value.hour) * 60 + cint(value.minute)
	text = str(value).strip()
	match = re.search(r"(\d{1,2}):(\d{2})", text)
	if not match:
		return None
	return cint(match.group(1)) * 60 + cint(match.group(2))


def _time_string_from_value(value, default=None):
	minutes = _minutes_from_time_value(value, default=default)
	if minutes is None:
		return None
	return f"{minutes // 60:02d}:{minutes % 60:02d}:00"


def _minutes_to_time_string(minutes):
	if minutes is None:
		return None
	minutes = int(minutes)
	return f"{minutes // 60:02d}:{minutes % 60:02d}:00"


def _load_attendance_params():
	"""读「薪资考勤参数」Single（含排班档案）。表缺失/字段空 → 回退默认，保证不崩。"""
	doc = None
	try:
		if frappe.db.exists("DocType", "Payroll Attendance Parameter"):
			doc = frappe.get_cached_doc("Payroll Attendance Parameter")
	except Exception:
		doc = None
	meal_price = _param_value(doc, "meal_unit_price", DEFAULT_MEAL_UNIT_PRICE, flt)
	rest_days = _param_value(doc, "non_worker_monthly_rest_days", NON_WORKER_MONTHLY_REST_DAYS, cint)
	threshold = _param_value(doc, "deduction_threshold_minutes", ATTENDANCE_DEDUCTION_THRESHOLD_MINUTES, cint)
	divisor = _param_value(doc, "regular_daily_divisor", DEFAULT_REGULAR_DAILY_DIVISOR, cint)
	paid_leave_names = _split_names(_param_value(doc, "paid_leave_names", ",".join(DEFAULT_PAID_LEAVE_NAMES)))
	overtime_start_time = _time_string_from_value(
		_param_value(doc, "overtime_start_time", DEFAULT_OT_START),
		default=DEFAULT_OT_START,
	)
	overtime_min_minutes = _param_value(doc, "overtime_min_minutes", DEFAULT_OT_MIN_MINUTES, cint)
	overtime_step_minutes = _param_value(doc, "overtime_step_minutes", DEFAULT_OT_STEP_MIN, cint)
	overtime_cap_hours = _param_value(doc, "overtime_cap_hours", DEFAULT_OT_CAP_HOURS, flt)
	overtime_hours_per_day = _param_value(doc, "overtime_hours_per_day", DEFAULT_OT_HOURS_PER_DAY, flt)
	overtime_daily_divisor = _param_value(doc, "overtime_daily_divisor", DEFAULT_OVERTIME_DAILY_DIVISOR, cint)
	profiles = {}
	for row in (doc.get("schedule_profiles") if doc else None) or []:
		key = "regular" if cint(row.is_regular) else "non_worker"
		profiles[key] = {
			"work_start": row.work_start,
			"lunch_start": row.lunch_start,
			"lunch_end": row.lunch_end,
			"work_end": row.work_end,
			"daily_hours": flt(row.daily_hours),
		}
	return frappe._dict(
		meal_unit_price=meal_price,
		non_worker_monthly_rest_days=rest_days,
		deduction_threshold_minutes=threshold,
		regular_daily_divisor=divisor,
		paid_leave_names=paid_leave_names or list(DEFAULT_PAID_LEAVE_NAMES),
		overtime_start_time=overtime_start_time,
		overtime_min_minutes=overtime_min_minutes,
		overtime_step_minutes=overtime_step_minutes,
		overtime_cap_hours=overtime_cap_hours,
		overtime_hours_per_day=overtime_hours_per_day,
		overtime_daily_divisor=overtime_daily_divisor,
		profiles=profiles,
	)


def _schedule_profile_for(is_regular):
	"""返回该岗位排班档案（分钟整数 + 日工时）。普工/综合岗各一套，缺失回退默认。"""
	key = "regular" if is_regular else "non_worker"
	raw = _load_attendance_params().profiles.get(key) or {}
	base = DEFAULT_SCHEDULE_PROFILES[key]

	def _m(field):
		return _minutes_from_time_value(raw.get(field), default=base[field])

	daily_hours = flt(raw.get("daily_hours")) or base["daily_hours"]
	return frappe._dict(
		key=key,
		work_start=_m("work_start"),
		lunch_start=_m("lunch_start"),
		lunch_end=_m("lunch_end"),
		work_end=_m("work_end"),
		daily_hours=daily_hours,
		daily_minutes=int(round(daily_hours * 60)),
	)


def _work_minutes_between(start_min, end_min, lunch_start, lunch_end):
	"""[start,end] 区间内扣除午休 overlap 后的工作分钟。"""
	if start_min is None or end_min is None or end_min <= start_min:
		return 0
	overlap = max(0, min(end_min, lunch_end) - max(start_min, lunch_start))
	return max(0, (end_min - start_min) - overlap)


def missed_whole_half(in_min, out_min, profile):
	"""缺整个上午(到岗≥午休起) 或 缺整个下午(离岗≤午休止) → True（用于餐补归 0）。"""
	if in_min is None or out_min is None:
		return True
	return in_min >= profile.lunch_start or out_min <= profile.lunch_end


def calculate_late_early_minutes(
	in_time=None, out_time=None, scheduled_in=None, scheduled_out=None,
	has_leave=False, profile=None, threshold=None,
):
	"""返回达阈值后的迟到/早退【缺勤工作分钟】（撇除午休），每天合计封顶日工时。

	迟到 = 上班时刻→实际上班 的工作分钟；早退 = 实际下班→下班时刻 的工作分钟。
	不足阈值记 0；满阈值含前段全计。普工/非普工按各自档案。"""
	if has_leave:
		return frappe._dict(late_minutes=0, early_minutes=0)
	if profile is None:
		profile = _schedule_profile_for(False)
	if threshold is None:
		threshold = _load_attendance_params().deduction_threshold_minutes

	work_start = _minutes_from_time_value(scheduled_in) or profile.work_start
	work_end = _minutes_from_time_value(scheduled_out) or profile.work_end
	in_minutes = _minutes_from_time_value(in_time)
	out_minutes = _minutes_from_time_value(out_time)

	late_minutes = (
		_work_minutes_between(work_start, in_minutes, profile.lunch_start, profile.lunch_end)
		if in_minutes is not None and in_minutes > work_start else 0
	)
	early_minutes = (
		_work_minutes_between(out_minutes, work_end, profile.lunch_start, profile.lunch_end)
		if out_minutes is not None and out_minutes < work_end else 0
	)
	if late_minutes < threshold:
		late_minutes = 0
	if early_minutes < threshold:
		early_minutes = 0
	cap = profile.daily_minutes
	if late_minutes + early_minutes > cap:
		early_minutes = max(0, cap - late_minutes)
		late_minutes = min(late_minutes, cap)
	return frappe._dict(late_minutes=late_minutes, early_minutes=early_minutes)


def _overtime_params():
	params = _load_attendance_params()
	step_minutes = cint(params.overtime_step_minutes) or DEFAULT_OT_STEP_MIN
	hours_per_day = flt(params.overtime_hours_per_day) or DEFAULT_OT_HOURS_PER_DAY
	return frappe._dict(
		start_minutes=_minutes_from_time_value(params.overtime_start_time, default=DEFAULT_OT_START) or _to_minutes(DEFAULT_OT_START),
		min_minutes=max(cint(params.overtime_min_minutes), 0),
		step_minutes=max(step_minutes, 1),
		cap_hours=max(flt(params.overtime_cap_hours), 0),
		hours_per_day=max(hours_per_day, 0.01),
		daily_divisor=cint(params.overtime_daily_divisor) or DEFAULT_OVERTIME_DAILY_DIVISOR,
	)


def overtime_days_for_regular(out_time, is_absent: bool = False) -> float:
	"""普工加班折天；参数来自「薪资考勤参数」，旷工 / 无下班卡 = 0。"""
	if is_absent or not out_time:
		return 0.0
	out_minutes = _minutes_from_time_value(out_time)
	if out_minutes is None:
		return 0.0
	params = _overtime_params()
	minutes = out_minutes - params.start_minutes
	if minutes < params.min_minutes:
		return 0.0
	steps = minutes // params.step_minutes
	hours = min(steps * (params.step_minutes / 60.0), params.cap_hours)
	return hours / params.hours_per_day


def daily_overtime_days(out_time, is_absent: bool, is_regular: bool) -> float:
	"""按岗位分流：普工走阶梯算法，非普工一律 0（调休处理，不计加班费）。"""
	if not is_regular:
		return 0.0
	return overtime_days_for_regular(out_time, is_absent)


def parse_day_cell(cell):
	"""解析钉钉每日格 '车间工人班次:正常\\n(07:18,11:30,12:28,20:31)'；空/非当月 → None。"""
	if cell is None or not str(cell).strip():
		return None
	text = str(cell).strip()
	status_line, _, punch_line = text.partition("\n")
	statuses = [seg.split(":")[-1] for seg in status_line.split(",")]
	is_absent = any("旷工" in s for s in statuses)
	is_leave = any(k in status_line for k in ("请假", "年假", "事假", "病假", "调休", "丧假"))
	is_late = "迟到" in status_line  # 含「严重迟到」
	is_early = "早退" in status_line
	m = re.search(r"\(([^)]*)\)", punch_line)
	punches = [p.strip() for p in m.group(1).split(",")] if m else []
	times = [p for p in punches if p and p != "-"]
	return frappe._dict(
		status_text=status_line,
		punch_summary=text.replace("\n", " "),
		in_time=(times[0] if times else None),
		out_time=(times[-1] if times else None),
		is_absent=is_absent,
		is_leave=is_leave,
		leave_name=_leave_name_from_text(status_line),
		is_late=is_late,
		is_early=is_early,
		overtime_days=0.0,  # 占位；真实值在 import_monthly_summary 按岗位(是否普工)覆盖
	)


def _split_names(value):
	text = str(value or "")
	return [item.strip() for item in re.split(r"[\s,，;；|]+", text) if item.strip()]


def _paid_leave_names(settings=None):
	params = _load_attendance_params()
	if params.paid_leave_names:
		return list(params.paid_leave_names)
	if settings and settings.get("paid_leave_names"):
		names = _split_names(settings.get("paid_leave_names"))
		if names:
			return names
	return list(DEFAULT_PAID_LEAVE_NAMES)


def is_paid_dingtalk_leave(leave_name, settings=None):
	text = str(leave_name or "")
	if not text:
		return False
	return any(name and name in text for name in _paid_leave_names(settings))


def _leave_name_from_text(text):
	text = str(text or "")
	for name in ("年假", "丧假", "事假", "病假", "调休", "请假"):
		if name in text:
			return name
	return ""


def parse_monthly_summary(rows):
	"""rows = list(ws.iter_rows(values_only=True))。返回 (start_date, end_date, [record])。
	record = _dict(name, user_id, days=[_dict(date, **parse_day_cell)])。"""
	title = " ".join(str(c) for c in rows[0] if c)
	m = re.search(r"(\d{4}-\d{2}-\d{2})\D+(\d{4}-\d{2}-\d{2})", title)
	start_date, end_date = (getdate(m.group(1)), getdate(m.group(2))) if m else (None, None)
	hdr = rows[2]

	def col(label):
		for i, c in enumerate(hdr):
			if c and label in str(c):
				return i
		return None

	name_c, uid_c, day_start = col("姓名"), col("UserId"), col("考勤结果")
	if name_c is None or day_start is None or start_date is None:
		frappe.throw("无法识别钉钉月度汇总表头（姓名/UserId/考勤结果/统计日期）")
	n_days = (end_date - start_date).days + 1
	records = []
	for r in rows[4:]:
		if name_c >= len(r) or not r[name_c]:
			continue
		days = []
		for k in range(n_days):
			ci = day_start + k
			parsed = parse_day_cell(r[ci] if ci < len(r) else None)
			if parsed:
				parsed.date = add_days(start_date, k)
				days.append(parsed)
		records.append(
			frappe._dict(
				name=str(r[name_c]).strip(),
				user_id=(str(r[uid_c]).strip() if uid_c is not None and uid_c < len(r) and r[uid_c] else ""),
				days=days,
			)
		)
	return start_date, end_date, records


# --------------------------------------------------------------------------- #
# 以下依赖 DB
# --------------------------------------------------------------------------- #
def _read_xlsx_rows(file_url: str):
	import openpyxl

	path = frappe.get_site_path(file_url.lstrip("/")) if file_url.startswith("/files") else frappe.utils.get_files_path(
		file_url.split("/")[-1]
	)
	wb = openpyxl.load_workbook(path, data_only=True)
	return list(wb.worksheets[0].iter_rows(values_only=True))


def _match_employee(user_id: str, name: str):
	if user_id:
		emp = frappe.db.get_value("Employee", {"attendance_device_id": user_id}, "name")
		if emp:
			return emp
	if name:
		return frappe.db.get_value("Employee", {"employee_name": name, "status": "Active"}, "name")
	return None


def _upsert_attendance(emp, date, status, parsed, company, is_regular=False):
	in_dt = f"{date} {parsed.in_time}:00" if parsed.in_time else None
	out_dt = f"{date} {parsed.out_time}:00" if parsed.out_time else None
	leave_days = 1.0 if parsed.get("is_leave") else 0.0
	paid_leave_days = leave_days if is_paid_dingtalk_leave(parsed.get("leave_name")) else 0.0
	unpaid_leave_days = leave_days - paid_leave_days
	# 单边卡（只上班 或 只下班）→ 缺勤（员工自行补卡）
	has_both = bool(parsed.in_time) and bool(parsed.out_time) and parsed.in_time != parsed.out_time
	if not parsed.get("is_leave") and not parsed.get("is_absent") and not has_both:
		status = "Absent"
	profile = _schedule_profile_for(is_regular)
	actual_attendance_days = 0.0 if parsed.get("is_leave") else (1.0 if has_both else 0.0)
	scheduled_in_time = _time_string_from_value(parsed.get("scheduled_in_time")) or _minutes_to_time_string(profile.work_start)
	scheduled_out_time = _time_string_from_value(parsed.get("scheduled_out_time")) or _minutes_to_time_string(profile.work_end)
	in_min = _minutes_from_time_value(in_dt)
	out_min = _minutes_from_time_value(out_dt)
	whole_half_missed = missed_whole_half(in_min, out_min, profile) if has_both else True
	meal_days = 1.0 if (actual_attendance_days == 1.0 and not leave_days and not whole_half_missed) else 0.0
	minute_factors = calculate_late_early_minutes(
		in_time=in_dt,
		out_time=out_dt,
		scheduled_in=scheduled_in_time,
		scheduled_out=scheduled_out_time,
		has_leave=bool(leave_days),
		profile=profile,
	)
	vals = {
		"status": status,
		"in_time": in_dt,
		"out_time": out_dt,
		"late_entry": 1 if parsed.get("is_late") or minute_factors.late_minutes else 0,
		"early_exit": 1 if parsed.get("is_early") or minute_factors.early_minutes else 0,
		"custom_punch_summary": parsed.punch_summary,
		"custom_overtime_days": parsed.overtime_days,
		"custom_dingtalk_leave_name": parsed.get("leave_name"),
		"custom_dingtalk_paid_leave_days": paid_leave_days,
		"custom_dingtalk_unpaid_leave_days": unpaid_leave_days,
		"custom_actual_attendance_days": actual_attendance_days,
		"custom_meal_allowance_days": meal_days,
		"custom_dingtalk_scheduled_in_time": scheduled_in_time,
		"custom_dingtalk_scheduled_out_time": scheduled_out_time,
		"custom_late_minutes": minute_factors.late_minutes,
		"custom_early_minutes": minute_factors.early_minutes,
	}
	existing = frappe.db.exists("Attendance", {"employee": emp, "attendance_date": str(date), "docstatus": ["<", 2]})
	if existing:
		frappe.db.set_value("Attendance", existing, vals, update_modified=False)
		return "updated"
	doc = frappe.get_doc(
		dict(doctype="Attendance", employee=emp, attendance_date=str(date), company=company, **vals)
	)
	doc.flags.ignore_validate = True
	doc.insert(ignore_permissions=True)
	doc.submit()
	return "created"


def _employee_structure(emp, on_date):
	"""员工在 on_date 生效的最新 SSA 工资结构名（与 get_overtime_amount 同口径）。无则 None。"""
	return frappe.db.get_value(
		"Salary Structure Assignment",
		{"employee": emp, "from_date": ("<=", on_date), "docstatus": 1},
		"salary_structure",
		order_by="from_date desc",
	)


@frappe.whitelist()
def import_monthly_summary(file_url=None, rows=None, company=None):
	"""从钉钉月度汇总建/更 Attendance（状态+in/out+打卡串+加班天）。返回统计。
	加班天仅对「普工结构」员工按阶梯算（非普工=0，调休处理）；无 SSA 的员工记入 no_structure。"""
	if rows is None:
		rows = _read_xlsx_rows(file_url)
	start_date, end_date, records = parse_monthly_summary(rows)
	stat = frappe._dict(
		start=str(start_date), end=str(end_date), created=0, updated=0,
		unmatched=[], no_structure=[],
	)
	for rec in records:
		emp = _match_employee(rec.user_id, rec.name)
		if not emp:
			stat.unmatched.append(rec.name)
			continue
		structure = _employee_structure(emp, start_date)
		if not structure:
			stat.no_structure.append(rec.name)
		is_regular = structure == REGULAR_WORKER_STRUCTURE
		emp_company = company or frappe.db.get_value("Employee", emp, "company")
		for d in rec.days:
			d.overtime_days = daily_overtime_days(d.out_time, d.is_absent, is_regular)
			status = "Absent" if d.is_absent else ("On Leave" if d.is_leave else "Present")
			stat[_upsert_attendance(emp, d.date, status, d, emp_company, is_regular)] += 1
	return stat


def get_overtime_amount(employee, start_date, end_date):
	"""返回 (加班天合计, 加班金额)。金额 = Σ加班天 × 总月薪/参数表分母。"""
	total_days = flt(
		frappe.db.sql(
			"""select coalesce(sum(custom_overtime_days),0) from `tabAttendance`
			where employee=%s and attendance_date between %s and %s and docstatus=1""",
			(employee, start_date, end_date),
		)[0][0]
	)
	if not total_days:
		return 0.0, 0.0
	ssa = frappe.db.get_value(
		"Salary Structure Assignment",
		{"employee": employee, "from_date": ("<=", start_date), "docstatus": 1},
		["base", "custom_seniority_base", "custom_grade_base", "custom_cert_base"],
		order_by="from_date desc",
		as_dict=True,
	)
	if not ssa:
		return total_days, 0.0
	monthly = flt(ssa.base) + flt(ssa.custom_seniority_base) + flt(ssa.custom_grade_base) + flt(ssa.custom_cert_base)
	divisor = _overtime_params().daily_divisor or DEFAULT_OVERTIME_DAILY_DIVISOR
	return total_days, total_days * monthly / divisor


def _period_year_months(start_date, end_date):
	start = getdate(start_date)
	end = getdate(end_date)
	year = start.year
	month = start.month
	while (year, month) <= (end.year, end.month):
		yield year, month
		if month == 12:
			year += 1
			month = 1
		else:
			month += 1


def _month_bounds(year, month, start_date, end_date):
	start = getdate(start_date)
	end = getdate(end_date)
	month_start = getdate(f"{year}-{month:02d}-01")
	next_month = getdate(f"{year + 1}-01-01") if month == 12 else getdate(f"{year}-{month + 1:02d}-01")
	month_end = add_days(next_month, -1)
	return max(start, month_start), min(end, month_end)


def _holiday_list_days(employee, start_date, end_date):
	holiday_list = frappe.db.get_value("Employee", employee, "holiday_list")
	if not holiday_list:
		return 0.0
	return flt(
		frappe.db.count(
			"Holiday",
			{
				"parent": holiday_list,
				"holiday_date": ["between", [start_date, end_date]],
			},
		)
	)


def _holiday_quota_year_rows(years):
	years = sorted({cint(year) for year in years if cint(year)})
	if not years:
		return []
	try:
		if not frappe.db.exists("DocType", "Payroll Holiday Quota Year"):
			return []
		placeholders = ", ".join(["%s"] * len(years))
		return frappe.db.sql(
			f"""
			select parent.year, child.holiday_month, child.holiday_name,
			       child.is_regular, child.quota_days, child.enabled
			from `tabPayroll Holiday Quota Year` parent
			inner join `tabPayroll Holiday Quota` child
			    on child.parent = parent.name
			   and child.parenttype = 'Payroll Holiday Quota Year'
			   and child.parentfield = 'holiday_quotas'
			where parent.year in ({placeholders})
			""",
			years,
			as_dict=True,
		)
	except Exception:
		return []


def _legacy_parameter_holiday_quota_rows():
	"""旧结构兼容：迁移前法定假配额曾挂在 Payroll Attendance Parameter 子表。"""
	try:
		return frappe.db.sql(
			"""
			select year, holiday_month, holiday_name, is_regular, quota_days, enabled
			from `tabPayroll Holiday Quota`
			where parenttype = 'Payroll Attendance Parameter'
			  and parentfield = 'holiday_quotas'
			""",
			as_dict=True,
		)
	except Exception:
		return []


def _holiday_quota_days(employee, start_date, end_date):
	"""法定假配额制。

	有该岗位+年份的年度配额单据时，按工资期间月份汇总配额；
	未配置的年份保留旧 Holiday List 计数兜底。
	"""
	structure = _employee_structure(employee, start_date)
	if not structure:
		return _holiday_list_days(employee, start_date, end_date)
	is_regular = structure == REGULAR_WORKER_STRUCTURE
	period_months = list(_period_year_months(start_date, end_date))
	rows = _holiday_quota_year_rows(year for year, _month in period_months)
	if not rows:
		rows = _legacy_parameter_holiday_quota_rows()
	group_rows = [
		row for row in rows
		if cint(row.get("enabled", 1)) and (1 if cint(row.get("is_regular")) else 0) == (1 if is_regular else 0)
	]
	if not group_rows:
		return _holiday_list_days(employee, start_date, end_date)

	configured_years = set()
	rows_by_year_month = {}
	for row in group_rows:
		year = cint(row.get("year"))
		month = cint(row.get("holiday_month"))
		if not year or month < 1 or month > 12:
			continue
		configured_years.add(year)
		rows_by_year_month.setdefault((year, month), 0.0)
		rows_by_year_month[(year, month)] += flt(row.get("quota_days"))

	total = 0.0
	for year, month in period_months:
		if year in configured_years:
			total += rows_by_year_month.get((year, month), 0.0)
		else:
			range_start, range_end = _month_bounds(year, month, start_date, end_date)
			total += _holiday_list_days(employee, range_start, range_end)
	return flt(total)


def _expected_work_days(employee, start_date, end_date, legal_holiday_days=None):
	natural_days = date_diff(end_date, start_date) + 1
	if legal_holiday_days is None:
		legal_holiday_days = _holiday_quota_days(employee, start_date, end_date)
	rest_days = _load_attendance_params().non_worker_monthly_rest_days
	return max(0.0, flt(natural_days) - flt(rest_days) - flt(legal_holiday_days))


def _attendance_factor_from_row(row):
	status = row.get("status")
	paid_leave_days = flt(row.get("custom_dingtalk_paid_leave_days"))
	unpaid_leave_days = flt(row.get("custom_dingtalk_unpaid_leave_days"))
	has_leave_days = bool(paid_leave_days or unpaid_leave_days)
	stored_actual_days = flt(row.get("custom_actual_attendance_days"))
	stored_meal_days = flt(row.get("custom_meal_allowance_days"))

	if stored_actual_days:
		actual_days = stored_actual_days
	elif status == "Present" and not has_leave_days:
		actual_days = 1.0
	elif status == "Half Day":
		actual_days = 0.5
	else:
		actual_days = 0.0

	# 餐补直接信任同步写入值（缺整半天=0 已在同步层判定）；不再按 Present 兜底为 1，避免覆盖正确的 0。
	meal_days = stored_meal_days

	return frappe._dict(
		actual_attendance_days=actual_days,
		paid_leave_days=paid_leave_days,
		unpaid_leave_days=unpaid_leave_days,
		meal_days=meal_days,
	)


def _attendance_minute_factors_from_row(row):
	"""信任同步层写入的缺勤工作分钟（已撇午休、按档案、封顶）。请假日不计。"""
	paid_leave_days = flt(row.get("custom_dingtalk_paid_leave_days"))
	unpaid_leave_days = flt(row.get("custom_dingtalk_unpaid_leave_days"))
	if paid_leave_days or unpaid_leave_days:
		return frappe._dict(late_minutes=0.0, early_minutes=0.0)
	return frappe._dict(
		late_minutes=flt(row.get("custom_late_minutes")),
		early_minutes=flt(row.get("custom_early_minutes")),
	)


def get_attendance_payroll_factors(employee, start_date, end_date):
	"""返回自定义工资天数与餐补天数。

	工资天数 = 实际到岗 + 钉钉带薪假 + 当月法定假数量。
	餐补天数 = 全天实际到岗；半天到岗/任何请假/法定假未到岗均为 0。
	"""
	rows = frappe.db.sql(
		"""select status,
		          in_time,
		          out_time,
		          custom_dingtalk_paid_leave_days,
		          custom_dingtalk_unpaid_leave_days,
		          custom_actual_attendance_days,
		          custom_meal_allowance_days,
		          custom_dingtalk_scheduled_in_time,
		          custom_dingtalk_scheduled_out_time,
		          custom_late_minutes,
		          custom_early_minutes
		   from `tabAttendance`
		   where employee=%s and attendance_date between %s and %s and docstatus=1""",
		(employee, start_date, end_date),
		as_dict=True,
	)
	actual_days = paid_leave_days = unpaid_leave_days = meal_days = 0.0
	late_minutes = early_minutes = 0.0
	for row in rows:
		factors = _attendance_factor_from_row(row)
		actual_days += factors.actual_attendance_days
		paid_leave_days += factors.paid_leave_days
		unpaid_leave_days += factors.unpaid_leave_days
		meal_days += factors.meal_days
		minute_factors = _attendance_minute_factors_from_row(row)
		late_minutes += minute_factors.late_minutes
		early_minutes += minute_factors.early_minutes

	legal_holiday_days = _holiday_quota_days(employee, start_date, end_date)
	expected_work_days = _expected_work_days(employee, start_date, end_date, legal_holiday_days)
	unpaid_absence_days = max(0.0, expected_work_days - actual_days - paid_leave_days)
	return frappe._dict(
		attendance_count=len(rows),
		actual_attendance_days=actual_days,
		paid_leave_days=paid_leave_days,
		unpaid_leave_days=unpaid_leave_days,
		legal_holiday_days=legal_holiday_days,
		salary_days=actual_days + paid_leave_days + legal_holiday_days,
		meal_days=meal_days,
		expected_work_days=expected_work_days,
		unpaid_absence_days=unpaid_absence_days,
		late_minutes=late_minutes,
		early_minutes=early_minutes,
	)


def _component_abbr(component):
	return frappe.db.get_value("Salary Component", component, "salary_component_abbr") or str(component)[:6]


def _set_component_amount(doc, tablefield, component, amount, preserve_zero=True, append_missing=False, freeze_formula=True):
	row = next((e for e in doc.get(tablefield, []) if e.salary_component == component), None)
	if row is None:
		if not append_missing or (preserve_zero and not amount) or not frappe.db.exists("Salary Component", component):
			return False
		row = doc.append(
			tablefield,
			{
				"salary_component": component,
				"abbr": _component_abbr(component),
				"amount": 0,
				"default_amount": 0,
			},
		)
	if preserve_zero and not amount:
		return False
	new_amt = flt(amount, row.precision("amount"))
	changed = flt(row.amount) != new_amt or flt(row.default_amount) != new_amt
	row.amount = row.default_amount = new_amt
	if freeze_formula:
		row.amount_based_on_formula = 0
		row.formula = None
		row.depends_on_payment_days = 0
	if changed:
		return True
	return False


def _remove_component_rows(doc, tablefield, component):
	rows = [row for row in doc.get(tablefield, []) if row.salary_component == component]
	for row in rows:
		doc.get(tablefield).remove(row)
	return bool(rows)


def _set_earning_component_amount(doc, component, amount, preserve_zero=True):
	return _set_component_amount(doc, "earnings", component, amount, preserve_zero=preserve_zero)


def _salary_structure_for_slip(doc):
	return doc.get("salary_structure") or _employee_structure(doc.employee, doc.start_date)


def _salary_structure_assignment_for_slip(doc, structure=None):
	filters = {"employee": doc.employee, "from_date": ("<=", doc.start_date), "docstatus": 1}
	if structure:
		filters["salary_structure"] = structure
	return frappe.db.get_value(
		"Salary Structure Assignment",
		filters,
		["base", "custom_seniority_base", "custom_grade_base", "custom_cert_base"],
		order_by="from_date desc",
		as_dict=True,
	)


def _fixed_monthly_amount(ssa):
	return sum(flt(ssa.get(fieldname)) for _component, fieldname in FIXED_ATTENDANCE_COMPONENTS)


def _employee_requires_personal_social_security(employee):
	return frappe.db.get_value("Employee", employee, "employment_type") == SOCIAL_SECURITY_EMPLOYMENT_TYPE


def _social_security_personal_amount_for_date(start_date):
	from fwb_utils.fwb_manufacturing.doctype.payroll_social_security_parameter.payroll_social_security_parameter import (
		get_personal_amount_for_date,
	)
	return flt(get_personal_amount_for_date(start_date))


def _apply_social_security_adjustment(doc):
	result = frappe._dict(personal_amount=0, warning="", touched=False)
	result.touched = _remove_component_rows(doc, "earnings", COMPANY_SOCIAL_SECURITY_COMPONENT)
	result.touched = _remove_component_rows(doc, "deductions", COMPANY_SOCIAL_SECURITY_COMPONENT) or result.touched

	if not doc.get("employee") or not doc.get("start_date"):
		return result

	if not _employee_requires_personal_social_security(doc.employee):
		result.touched = _remove_component_rows(doc, "deductions", PERSONAL_SOCIAL_SECURITY_COMPONENT) or result.touched
		return result

	amount = _social_security_personal_amount_for_date(doc.start_date)
	if amount <= 0:
		result.touched = _remove_component_rows(doc, "deductions", PERSONAL_SOCIAL_SECURITY_COMPONENT) or result.touched
		result.warning = "当月未配置社保个人扣款，未扣社保。"
		return result

	if not frappe.db.exists("Salary Component", PERSONAL_SOCIAL_SECURITY_COMPONENT):
		result.warning = "工资构成「社保-个人」不存在，未扣社保。"
		return result

	result.touched = (
		_set_component_amount(
			doc,
			"deductions",
			PERSONAL_SOCIAL_SECURITY_COMPONENT,
			amount,
			preserve_zero=False,
			append_missing=True,
		)
		or result.touched
	)
	result.personal_amount = amount
	return result


def _apply_non_worker_salary_adjustments(doc, factors, structure):
	ssa = _salary_structure_assignment_for_slip(doc, structure)
	if not ssa:
		return False

	expected_work_days = max(flt(factors.expected_work_days), 1.0)
	work_ratio = max(0.0, expected_work_days - flt(factors.unpaid_absence_days)) / expected_work_days
	monthly_fixed = _fixed_monthly_amount(ssa)
	touched = False

	for component, fieldname in FIXED_ATTENDANCE_COMPONENTS:
		component_base = flt(ssa.get(fieldname))
		touched = (
			_set_component_amount(
				doc,
				"earnings",
				component,
				component_base * work_ratio,
				preserve_zero=False,
				append_missing=bool(component_base),
			)
			or touched
		)

	meal_amount = flt(factors.meal_days) * flt(_load_attendance_params().meal_unit_price)
	touched = _set_component_amount(doc, "earnings", MEAL_COMPONENT, meal_amount, preserve_zero=False) or touched

	daily_hours = _schedule_profile_for(False).daily_hours or NON_WORKER_STANDARD_HOURS_PER_DAY
	minute_rate = monthly_fixed / expected_work_days / daily_hours / 60.0 if monthly_fixed else 0.0
	late_amount = flt(factors.late_minutes) * minute_rate
	early_amount = flt(factors.early_minutes) * minute_rate
	touched = (
		_set_component_amount(
			doc,
			"deductions",
			LATE_DEDUCTION_COMPONENT,
			late_amount,
			preserve_zero=False,
			append_missing=True,
		)
		or touched
	)
	touched = (
		_set_component_amount(
			doc,
			"deductions",
			EARLY_DEDUCTION_COMPONENT,
			early_amount,
			preserve_zero=False,
			append_missing=True,
		)
		or touched
	)
	touched = _set_component_amount(doc, "earnings", OVERTIME_COMPONENT, 0, preserve_zero=False) or touched
	factors.fixed_monthly_amount = monthly_fixed
	factors.late_deduction_amount = late_amount
	factors.early_deduction_amount = early_amount
	return touched


def _apply_regular_worker_minute_deductions(doc, factors, structure):
	"""普工迟到/早退扣款：单价 = 月固定(底薪+工龄+等级)/单日工资分母/普工日工时/60；封顶已在同步层按日处理。

	普工口径独立于非普工：分母用「普工单日工资分母」(默认30)、普工日工时(默认8.5)，切勿与非普工的/应出勤混用。
	普工无餐补；payment_days 与底薪公式已按到岗折算整日缺勤，迟到/早退只走分钟扣款。"""
	ssa = _salary_structure_assignment_for_slip(doc, structure)
	if not ssa:
		return False
	params = _load_attendance_params()
	profile = _schedule_profile_for(True)
	monthly_fixed = sum(flt(ssa.get(fieldname)) for _component, fieldname in REGULAR_FIXED_COMPONENTS)
	divisor = cint(params.regular_daily_divisor) or DEFAULT_REGULAR_DAILY_DIVISOR
	daily_hours = profile.daily_hours or NON_WORKER_STANDARD_HOURS_PER_DAY
	minute_rate = monthly_fixed / divisor / daily_hours / 60.0 if monthly_fixed else 0.0
	late_amount = flt(factors.late_minutes) * minute_rate
	early_amount = flt(factors.early_minutes) * minute_rate
	touched = _set_component_amount(
		doc, "deductions", LATE_DEDUCTION_COMPONENT, late_amount, preserve_zero=False, append_missing=True
	)
	touched = (
		_set_component_amount(
			doc, "deductions", EARLY_DEDUCTION_COMPONENT, early_amount, preserve_zero=False, append_missing=True
		)
		or touched
	)
	factors.late_deduction_amount = late_amount
	factors.early_deduction_amount = early_amount
	return touched


def _refresh_salary_totals(doc):
	doc.gross_pay = doc.get_component_totals("earnings", depends_on_payment_days=1)
	doc.base_gross_pay = flt(flt(doc.gross_pay) * flt(doc.exchange_rate or 1), doc.precision("base_gross_pay"))
	doc.set_net_pay()


def apply_attendance_payroll_adjustments_to_salary_slip(doc, method=None):
	"""工资单按钮调用：统一按 Attendance 修正工资天数、餐补和加班。

	- payment_days 改为“实际到岗 + 带薪假 + 当月法定假额度”。
	- 餐补 = 全天实际到岗天数 × 参数表餐补单价；半天到岗/请假/法定假未到岗均为 0。
	- 加班按参数表普工加班规则计算；仅当有加班金额时覆盖「加班」行。
	"""
	if doc.docstatus == 2 or not doc.get("employee") or not doc.get("start_date"):
		return

	factors = get_attendance_payroll_factors(doc.employee, doc.start_date, doc.end_date)
	if cint(factors.attendance_count):
		custom_payment_days = flt(factors.salary_days, doc.precision("payment_days"))
		if flt(doc.payment_days) != custom_payment_days:
			doc.payment_days = custom_payment_days
			doc.calculate_net_pay()

		structure = _salary_structure_for_slip(doc)
		if structure in NON_WORKER_STRUCTURES:
			_apply_non_worker_salary_adjustments(doc, factors, structure)
		elif structure == REGULAR_WORKER_STRUCTURE:
			_apply_regular_worker_minute_deductions(doc, factors, structure)
		else:
			meal_amount = flt(factors.meal_days) * flt(_load_attendance_params().meal_unit_price)
			_set_earning_component_amount(doc, MEAL_COMPONENT, meal_amount, preserve_zero=False)

	_, amount = get_overtime_amount(doc.employee, doc.start_date, doc.end_date)
	if _salary_structure_for_slip(doc) in NON_WORKER_STRUCTURES:
		amount = 0
	touched_overtime = _set_earning_component_amount(doc, OVERTIME_COMPONENT, amount, preserve_zero=True)
	social_security = _apply_social_security_adjustment(doc)
	if cint(factors.attendance_count) or touched_overtime or social_security.touched:
		_refresh_salary_totals(doc)
		doc.compute_year_to_date()
		doc.compute_month_to_date()
		doc.compute_component_wise_year_to_date()
	return frappe._dict(
		social_security_personal_amount=social_security.personal_amount,
		social_security_warning=social_security.warning,
	)


def apply_overtime_to_salary_slip(doc, method=None):
	"""兼容旧 hook 名称。实际逻辑已合并到 apply_attendance_payroll_adjustments_to_salary_slip。"""
	return apply_attendance_payroll_adjustments_to_salary_slip(doc, method=method)
